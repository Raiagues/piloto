#!/usr/bin/env python3
"""Build DOCX copies (same turn IDs) of the synthetic LYLE meetings and a short
SYNTHETIC audio with equivalent content (Piper TTS, local). Audio goes to .runtime
(outside Git) and is labelled synthetic in its file name and metadata."""
import json
import re
import sys
import wave
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
INPUTS = ROOT / 'norte' / 'examples' / 'lyle-synthetic' / 'inputs'
AUDIO_DIR = ROOT / '.runtime' / 'synthetic-audio'
HEADER = re.compile(r'^\s*([A-Z0-9]{1,8}-T\d{1,4})\s*\|\s*(.+?)\s*$')


def turns(path):
    out, current = [], None
    for line in path.read_text(encoding='utf-8').splitlines():
        match = HEADER.match(line)
        if match:
            current = {'id': match.group(1), 'speaker': match.group(2), 'text': ''}
            out.append(current)
        elif current and line.strip():
            current['text'] += (' ' if current['text'] else '') + line.strip()
    preamble = path.read_text(encoding='utf-8').split('\n', 1)[0]
    return preamble, out


def build_docx(name):
    import docx
    preamble, items = turns(INPUTS / f'{name}.txt')
    document = docx.Document()
    document.add_paragraph(preamble)
    for turn in items:
        document.add_paragraph(f"{turn['id']} | {turn['speaker']}")
        document.add_paragraph(turn['text'])
    document.core_properties.title = name
    document.core_properties.comments = 'Conteúdo sintético de teste; pessoas fictícias.'
    target = INPUTS / f'{name}.docx'
    document.save(target)
    return target


def build_audio(name, max_turns=None):
    from piper import PiperVoice, SynthesisConfig
    voice = PiperVoice.load(ROOT / '.runtime' / 'models' / 'piper' / 'pt_BR-faber-medium.onnx')
    _, items = turns(INPUTS / f'{name}.txt')
    items = items[:max_turns] if max_turns else items
    AUDIO_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
    target = AUDIO_DIR / f'{name}-SINTETICO-piper.wav'
    rate = voice.config.sample_rate
    silence = b'\x00\x00' * int(rate * 0.6)
    segments = []
    with wave.open(str(target), 'wb') as out:
        out.setnchannels(1); out.setsampwidth(2); out.setframerate(rate)
        position = 0.0
        for turn in items:
            chunk_path = AUDIO_DIR / '.chunk.wav'
            spoken = re.sub(r'REQ-LYLE-0*(\d+)', r'requisito LYLE \1', turn['text'])  # IDs as people say them
            with wave.open(str(chunk_path), 'wb') as chunk:
                voice.synthesize_wav(spoken, chunk, syn_config=SynthesisConfig(length_scale=1.05))
            with wave.open(str(chunk_path), 'rb') as chunk:
                frames = chunk.readframes(chunk.getnframes())
                duration = chunk.getnframes() / chunk.getframerate()
            out.writeframes(frames); out.writeframes(silence)
            segments.append({'turnId': turn['id'], 'start': round(position, 3), 'end': round(position + duration, 3), 'text': turn['text'], 'spoken': spoken})
            position += duration + 0.6
        chunk_path.unlink(missing_ok=True)
    meta = {'synthetic': True, 'engine': 'piper-tts 1.8.0', 'voice': 'pt_BR-faber-medium', 'source': f'{name}.txt', 'note': 'Áudio sintético (TTS local) com conteúdo equivalente; uma única voz, sem identificação de falantes. Não é gravação real.', 'turns': segments, 'durationSeconds': round(position, 3)}
    (AUDIO_DIR / f'{name}-SINTETICO-piper.json').write_text(json.dumps(meta, ensure_ascii=False, indent=1))
    return target, meta


if __name__ == '__main__':
    for name in ('M1-rastreabilidade', 'M2-confidencial-simulado'):
        print('docx', build_docx(name).relative_to(ROOT))
    if '--audio' in sys.argv:
        path, meta = build_audio('M1-rastreabilidade')
        print('audio', path.relative_to(ROOT), meta['durationSeconds'], 's')
        path, meta = build_audio('M1-variacao-pauta-parcial')
        print('audio', path.relative_to(ROOT), meta['durationSeconds'], 's')
